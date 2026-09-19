"""No live database/model: exercise business guards with explicit repository doubles."""
import copy
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock
import pytest
from fastapi import HTTPException
from yuxi.services.changwei_service import ChangweiService
from yuxi.repositories.changwei_repository import ChangweiRepository


@pytest.fixture
def setup_service(monkeypatch):
    project = SimpleNamespace(id='p', owner_uid='owner', members={'member': 'member'})
    task = SimpleNamespace(id='t', project_id='p', creator_uid='owner', revision=3, status='confirmed',
        content={'material_ids': ['a'], 'modules': [{'id': '0', 'name': '施工情况', 'text': '原记录',
            'assignee': 'member', 'confirmed_by': 'member', 'confirmed_at': 'yesterday', 'sources': []}]})
    db = SimpleNamespace(commit=AsyncMock(), rollback=AsyncMock(), add=Mock())
    service = ChangweiService(db, 'owner')
    service.repo = SimpleNamespace(task=AsyncMock(return_value=task), project=AsyncMock(return_value=project),
        materials=AsyncMock(return_value=[SimpleNamespace(id='a')]), audit=Mock())
    monkeypatch.setattr('yuxi.services.changwei_service.view', lambda row: vars(row))
    return service, task, project


@pytest.mark.asyncio
async def test_old_revision_cannot_change_module(setup_service):
    service, task, _ = setup_service
    before = copy.deepcopy(task.content)
    with pytest.raises(HTTPException) as err:
        await service.update_module('t', '0', 2, '覆盖', True)
    assert err.value.status_code == 409
    assert task.content == before
    service.db.commit.assert_not_awaited()


@pytest.mark.asyncio
async def test_member_cannot_write_another_assignees_module(setup_service):
    service, task, _ = setup_service
    service.uid = 'member'
    task.content['modules'][0]['assignee'] = 'owner'
    with pytest.raises(HTTPException) as err:
        await service.update_module('t', '0', 3, '覆盖', True)
    assert err.value.status_code == 403
    service.db.commit.assert_not_awaited()


@pytest.mark.asyncio
async def test_edit_revokes_confirmation_and_increments_revision(setup_service):
    service, task, _ = setup_service
    await service.update_module('t', '0', 3, ' 新记录 ', False)
    assert task.content['modules'][0]['confirmed_by'] is None
    assert task.content['modules'][0]['confirmed_at'] is None
    assert task.content['modules'][0]['text'] == '新记录'
    assert task.status == 'draft' and task.revision == 4


@pytest.mark.asyncio
async def test_reassignment_revokes_confirmation(setup_service):
    service, task, _ = setup_service
    await service.update_module('t', '0', 3, '原记录', True, 'owner')
    assert task.content['modules'][0]['assignee'] == 'owner'
    assert task.content['modules'][0]['confirmed_by'] is None


@pytest.mark.asyncio
async def test_cannot_confirm_empty_text(setup_service):
    service, _, _ = setup_service
    with pytest.raises(HTTPException) as err:
        await service.update_module('t', '0', 3, ' \n ', True)
    assert err.value.status_code == 422


@pytest.mark.asyncio
async def test_selecting_foreign_material_is_rejected(setup_service):
    service, task, _ = setup_service
    with pytest.raises(HTTPException) as err:
        await service.select_materials('t', 3, ['foreign'])
    assert err.value.status_code == 422
    assert task.content['material_ids'] == ['a']


@pytest.mark.asyncio
async def test_selecting_materials_revokes_all_confirmations(setup_service):
    service, task, _ = setup_service
    await service.select_materials('t', 3, [])
    assert task.content['modules'][0]['confirmed_by'] is None
    assert task.status == 'draft' and task.revision == 4


@pytest.mark.asyncio
async def test_unconfirmed_module_cannot_generate_artifact(setup_service):
    service, task, _ = setup_service
    task.content['modules'][0]['confirmed_by'] = None
    with pytest.raises(HTTPException) as err:
        await service.generate('t', 3)
    assert err.value.status_code == 422
    service.db.add.assert_not_called()


@pytest.mark.asyncio
@pytest.mark.parametrize('uid,manage,status', [('outsider', False, 404), ('member', True, 403)])
async def test_repository_project_permission(uid, manage, status):
    project = SimpleNamespace(owner_uid='owner', members={'member': 'member'})
    repo = ChangweiRepository(SimpleNamespace(scalar=AsyncMock(return_value=project)), uid)
    with pytest.raises(HTTPException) as err:
        await repo.project('p', manage=manage)
    assert err.value.status_code == status


@pytest.mark.asyncio
async def test_corrupt_upload_rolls_back_without_commit(setup_service):
    service, _, _ = setup_service
    with pytest.raises(HTTPException) as err:
        await service.upload('p', 'bad.zip', b'not a zip')
    assert err.value.status_code == 422
    service.db.rollback.assert_awaited_once()
    service.db.commit.assert_not_awaited()


@pytest.mark.asyncio
async def test_exported_snapshot_remains_unchanged_after_edit(setup_service):
    import io
    from docx import Document
    service, task, project = setup_service
    task.kind, task.title, task.period = 'supervision_log', '日报', '2026-09-19'
    project.name, project.lot = '工程', '一标'
    await service.generate('t', 3)
    artifact = service.db.add.call_args.args[0]
    snapshot = bytes(artifact.data)
    await service.update_module('t', '0', 3, '修改后的记录', False)
    assert artifact.data == snapshot
    text = '\n'.join(p.text for p in Document(io.BytesIO(snapshot)).paragraphs)
    assert '原记录' in text
    assert '修改后的记录' not in text
    assert artifact.revision == 3 and task.revision == 4


@pytest.mark.asyncio
async def test_mixed_zip_is_atomic_when_second_file_is_corrupt(setup_service):
    import io
    from zipfile import ZipFile
    service, _, _ = setup_service
    stream = io.BytesIO()
    with ZipFile(stream, 'w') as archive:
        archive.writestr('valid.txt', '有效资料')
        archive.writestr('broken.docx', 'not an office archive')
    with pytest.raises(HTTPException) as err:
        await service.upload('p', 'mixed.zip', stream.getvalue())
    assert err.value.status_code == 422
    service.db.rollback.assert_awaited_once()
    service.db.commit.assert_not_awaited()


@pytest.mark.asyncio
async def test_corrupt_deflate_upload_returns_validation_error(setup_service):
    import io
    import struct
    from zipfile import ZipFile, ZIP_DEFLATED
    service, _, _ = setup_service
    stream = io.BytesIO()
    with ZipFile(stream, 'w', compression=ZIP_DEFLATED) as archive:
        archive.writestr('notes.txt', '施工记录' * 100)
    raw = bytearray(stream.getvalue())
    name_len, extra_len = struct.unpack_from('<HH', raw, 26)
    start = 30 + name_len + extra_len
    raw[start] = 0xff  # Reserved DEFLATE block type, corrupt compressed payload.
    with pytest.raises(HTTPException) as err:
        await service.upload('p', 'corrupt.zip', bytes(raw))
    assert err.value.status_code == 422
    service.db.rollback.assert_awaited_once()
    service.db.commit.assert_not_awaited()


@pytest.mark.asyncio
async def test_corrupt_office_xml_returns_validation_error(setup_service):
    import io
    from zipfile import ZipFile
    from docx import Document
    service, _, _ = setup_service
    good = io.BytesIO()
    Document().save(good)
    corrupt = io.BytesIO()
    with ZipFile(io.BytesIO(good.getvalue())) as source, ZipFile(corrupt, 'w') as target:
        for item in source.infolist():
            target.writestr(item.filename, b'<broken' if item.filename == 'word/document.xml' else source.read(item))
    with pytest.raises(HTTPException) as err:
        await service.upload('p', 'corrupt.docx', corrupt.getvalue())
    assert err.value.status_code == 422
    service.db.rollback.assert_awaited_once()
    service.db.commit.assert_not_awaited()


@pytest.mark.asyncio
async def test_ocr_merges_latest_locked_material_without_losing_other_pages(setup_service, monkeypatch):
    service, _, _ = setup_service
    stale = SimpleNamespace(filename='scan.pdf', data=b'pdf', content={
        'segments': [], 'warnings': ['第 1 页需要 OCR 或人工补录']})
    latest = SimpleNamespace(id='m', project_id='p', confirmed=True, parse_status='needs_attention', content={
        'segments': [{'location': '第 6 页', 'text': '另一请求刚写入的识别结果'}],
        'warnings': ['第 1 页需要 OCR 或人工补录', '第 7 页需要 OCR 或人工补录']})
    async def material(_id, lock=False):
        return latest if lock else stale
    service.repo.material = AsyncMock(side_effect=material)
    monkeypatch.setattr('yuxi.services.changwei_service.ocr_pdf_pages',
        lambda data, pages: [{'location': '第 1 页', 'text': '本请求识别结果', 'ocr': True}])
    await service.ocr_material('m')
    assert {s['location'] for s in latest.content['segments']} == {'第 1 页', '第 6 页'}
    assert latest.content['warnings'] == ['第 7 页需要 OCR 或人工补录']
    assert latest.confirmed is False
    assert latest.parse_status == 'needs_attention'


@pytest.mark.asyncio
async def test_material_write_query_locks_and_refreshes_identity_map():
    from sqlalchemy.dialects import postgresql
    material = SimpleNamespace(project_id='p')
    db = SimpleNamespace(scalar=AsyncMock(return_value=material))
    repo = ChangweiRepository(db, 'owner')
    repo.project = AsyncMock()
    assert await repo.material('m', lock=True) is material
    stmt = db.scalar.call_args.args[0]
    assert 'FOR UPDATE' in str(stmt.compile(dialect=postgresql.dialect()))
    assert stmt.get_execution_options()['populate_existing'] is True
