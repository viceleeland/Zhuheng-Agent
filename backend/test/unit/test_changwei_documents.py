"""资料解析的独立结果断言与负向回归。"""
import io
from zipfile import ZipFile
import pytest
from docx import Document
from openpyxl import Workbook
from yuxi.services.changwei_documents import unpack, parse_document, export_docx, normalize_station


def test_zip_rejects_traversal_and_nested_archive():
    """资料包路径越界和嵌套压缩必须拒绝。"""
    for name in ['../secret.txt', '/absolute.txt', 'nested.zip']:
        stream = io.BytesIO()
        with ZipFile(stream, 'w') as archive:
            archive.writestr(name, 'test')
        with pytest.raises(ValueError):
            unpack('materials.zip', stream.getvalue())


def test_excel_reports_uncached_formula_and_coordinates():
    """缺失公式缓存不能被当成零金额。"""
    book = Workbook()
    book.active.title = '资金台账'
    book.active.append(['单位：万元', '本期', '累计'])
    book.active.append(['工程款', 12.5, '=B2+10'])
    stream = io.BytesIO()
    book.save(stream)
    parsed = parse_document('台账.xlsx', stream.getvalue())
    assert parsed['status'] == 'needs_attention'
    assert '资金台账!C2' in parsed['warnings'][0]
    assert 'B2=12.5' in parsed['segments'][1]['text']
    assert '[公式待重算]' in parsed['segments'][1]['text']


def test_export_keeps_confirmed_text_and_source():
    """重新读取实际 Word 字节而非断言函数调用次数。"""
    data = export_docx('监理日志', '测试工程', '2026-05-25', [
        {'name': '施工情况', 'text': 'FHDL1+194 至 FHDL1+206，现场 12 人。', 'sources': ['原始记录 / 段落 2']}
    ], 3)
    paragraphs = [p.text for p in Document(io.BytesIO(data)).paragraphs]
    assert 'FHDL1+194 至 FHDL1+206，现场 12 人。' in paragraphs
    assert '资料依据：原始记录 / 段落 2' in paragraphs
    assert any('版本：3' in p for p in paragraphs)


def test_scanned_pdf_is_not_reported_as_ready():
    """图像页没有文字时必须请求 OCR。"""
    import pypdfium2 as pdfium
    stream = io.BytesIO()
    with pdfium.PdfDocument.new() as pdf:
        pdf.new_page(595, 842).close()
        pdf.save(stream)
    parsed = parse_document('扫描方案.pdf', stream.getvalue())
    assert parsed['status'] == 'needs_attention'
    assert not parsed['segments']
    assert '第 1 页需要 OCR' in parsed['warnings'][0]


def test_station_normalization_preserves_numbers():
    """桩号只规范写法，不猜测数值。"""
    assert normalize_station('左洞1加194到1加206，12人') == 'FHDL1+194到1+206，12人'


def test_excel_sparse_realistic_ledger_keeps_coordinates():
    book = Workbook()
    sheet = book.active
    sheet.title = '资金台账'
    sheet.merge_cells('A1:D1')
    sheet['A1'] = '单位：万元'
    sheet['A3'] = '工程款'
    sheet['D3'] = 12.5
    sheet['B5'] = 0
    stream = io.BytesIO()
    book.save(stream)
    parsed = parse_document('资金.xlsx', stream.getvalue())
    assert parsed['status'] == 'ready'
    assert [s['location'] for s in parsed['segments']] == ['资金台账!1', '资金台账!3', '资金台账!5']
    assert parsed['segments'][1]['text'] == 'A3=工程款 | D3=12.5'
    assert parsed['segments'][2]['text'] == 'B5=0'


def test_long_text_is_lossless_and_has_bounded_segments():
    text = '工程施工现场记录。' * 2000
    parsed = parse_document('记录.txt', text.encode('utf-8'))
    assert ''.join(s['text'] for s in parsed['segments']) == text
    assert all(len(s['text']) <= 4000 for s in parsed['segments'])
    assert parsed['segments'][0]['location'] == '正文 / 字符 1-4000'


def test_oversized_extracted_text_rejected():
    with pytest.raises(ValueError, match='200 万'):
        parse_document('large.txt', b'a' * 2_000_001)


def test_zip_file_count_limit():
    stream = io.BytesIO()
    with ZipFile(stream, 'w') as archive:
        for index in range(101):
            archive.writestr(f'{index}.txt', 'x')
    with pytest.raises(ValueError, match='100'):
        unpack('many.zip', stream.getvalue())


def test_corrupt_zip_rejected():
    from zipfile import BadZipFile
    with pytest.raises(BadZipFile):
        unpack('bad.zip', b'not-a-zip')


def test_scan_over_200_pages_preserves_pending_page_numbers():
    """Every scanned page must remain discoverable by the batch OCR action."""
    import pypdfium2 as pdfium
    stream = io.BytesIO()
    with pdfium.PdfDocument.new() as pdf:
        for _ in range(201):
            pdf.new_page(595, 842).close()
        pdf.save(stream)
    parsed = parse_document('long-scan.pdf', stream.getvalue())
    assert any('第 201 页需要 OCR' in warning for warning in parsed['warnings'])


def test_pdf_parse_and_ocr_do_not_overlap_pdfium_lifetimes(monkeypatch):
    """Exercise both worker entry points without risking a native PDFium crash."""
    import sys
    import time
    import threading
    from types import SimpleNamespace
    from concurrent.futures import ThreadPoolExecutor
    import pypdfium2
    from yuxi.services.changwei_documents import ocr_pdf_pages
    state = {'active': 0, 'maximum': 0, 'closed_pages': 0}
    counter_lock = threading.Lock()
    class Page:
        def get_textpage(self): return self
        def get_text_bounded(self): return 'verified document text with enough characters'
        def get_size(self): return (100, 100)
        def render(self, **kwargs): return self
        def to_pil(self): return object()
        def close(self): state['closed_pages'] += 1
    class Document:
        def __init__(self, data):
            with counter_lock:
                state['active'] += 1
                state['maximum'] = max(state['maximum'], state['active'])
            time.sleep(0.02)
        def __enter__(self): return self
        def __exit__(self, *args):
            time.sleep(0.01)
            with counter_lock: state['active'] -= 1
        def __len__(self): return 1
        def __getitem__(self, number): return Page()
    monkeypatch.setattr(pypdfium2, 'PdfDocument', Document)
    factory = SimpleNamespace(DocumentProcessorFactory=SimpleNamespace(
        get_processor=lambda name: SimpleNamespace(process_image=lambda image: 'OCR text')))
    monkeypatch.setitem(sys.modules, 'yuxi.knowledge.parser.factory', factory)
    with ThreadPoolExecutor(max_workers=4) as pool:
        jobs = [pool.submit(parse_document, 'a.pdf', b'pdf'), pool.submit(ocr_pdf_pages, b'pdf', [1]),
                pool.submit(parse_document, 'b.pdf', b'pdf'), pool.submit(ocr_pdf_pages, b'pdf', [1])]
        for job in jobs: job.result()
    assert state['maximum'] == 1
    assert state['active'] == 0
    assert state['closed_pages'] == 8
