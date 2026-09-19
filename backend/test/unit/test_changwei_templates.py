from io import BytesIO
import re

import pytest
from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.shared import Pt
from yuxi.services.changwei_templates import export_supervision_template


def template():
    doc = Document()
    doc.add_paragraph('监 理 日 志')
    doc.add_paragraph('填写人：旧填写人 日期：旧日期 白班')
    doc.add_paragraph('说明：本段来自原模板，生成时不得删除。')
    doc.add_paragraph('填写完成后按原流程归档。')
    table = doc.add_table(rows=7, cols=8)
    table.style = 'Table Grid'
    for index, text in enumerate(['天气', '旧天气', '气温', '旧气温', '风力', '旧风力', '风向', '旧风向']):
        table.cell(0, index).text = text
    for index, text in enumerate(['施工部位施工内容资源投入', '承包人质量安全', '监理巡视检验', '问题处理落实', '监理签发', '其他事项'], 1):
        table.cell(index, 0).text = text
        cell = table.cell(index, 1).merge(table.cell(index, 7))
        cell.text = '旧现场内容'
        cell.paragraphs[0].runs[0].font.size = Pt(12)
        cell.add_paragraph('旧附加段落')
        table.rows[index].height = Pt(30)
        table.rows[index].height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def test_template_replaces_old_values_and_preserves_style():
    original = template()
    long_text = '质量检查完成。' * 2000
    result = export_supervision_template(original, [
        {'name': '施工部位及施工内容', 'text': '新施工内容'},
        {'name': '施工形象及资源投入', 'text': '新资源'},
        {'name': '承包人质量检验和安全作业', 'text': long_text},
        {'name': '天气信息', 'text': '晴，气温二十度'},
    ], '2026-09-19', '本次填写人')
    doc = Document(BytesIO(result))
    all_text = '\n'.join(p.text for p in doc.paragraphs) + '\n'.join(c.text for r in doc.tables[0].rows for c in r.cells)
    assert '旧' not in all_text
    assert '白班' in doc.paragraphs[1].text
    assert '2026年9月19日' in doc.paragraphs[1].text
    assert '本次填写人' in doc.paragraphs[1].text
    assert doc.tables[0].cell(1, 1).text == '新施工内容\n新资源'
    assert doc.tables[0].cell(2, 1).text == long_text
    assert doc.tables[0].cell(2, 1).paragraphs[0].runs[0].font.size == Pt(12)
    assert doc.tables[0].rows[2].height_rule == WD_ROW_HEIGHT_RULE.EXACTLY
    assert doc.tables[0].style.name == 'Table Grid'
    assert doc.tables[0].cell(0, 1).text == '晴，气温二十度'
    assert doc.sections[0].page_width == Document(BytesIO(original)).sections[0].page_width
    assert doc.paragraphs[2].text == Document(BytesIO(original)).paragraphs[2].text


def test_structured_weather_replaces_every_value():
    result = export_supervision_template(template(), [], '', '', {'weather': '雨', 'temperature': '18', 'wind_force': '3', 'wind_direction': '东'})
    row = Document(BytesIO(result)).tables[0].rows[0]
    assert [row.cells[i].text for i in (1, 3, 5, 7)] == ['雨', '18', '3', '东']


def test_weather_module_text_is_split_into_original_template_columns():
    result = export_supervision_template(template(), [{
        'name': '天气信息',
        'text': '天气：晴；气温：26 ℃；风力：2 级；风向：东风。\n来源：和风天气。',
    }], '2026-09-19', '张三')
    row = Document(BytesIO(result)).tables[0].rows[0]
    assert [row.cells[i].text for i in (1, 3, 5, 7)] == ['晴', '26 ℃', '2 级', '东风']


def test_placeholder_weather_markdown_does_not_pollute_compact_cells():
    result = export_supervision_template(template(), [{
        'name': '天气信息',
        'text': '# 天气信息\n**天气（晴/阴/雨/雪等）：**【待补充】\n**气温（最高/最低，℃）：**【待补充】\n**风力风向：**【待补充】',
    }], '2026-09-19', '张三')
    row = Document(BytesIO(result)).tables[0].rows[0]
    assert [row.cells[i].text for i in (1, 3, 5, 7)] == ['', '', '', '']


def test_blank_header_placeholders_keep_shift_and_class_text():
    doc = Document(BytesIO(template()))
    doc.paragraphs[1].text = '填写人：          日期：    年  月  日     白班'
    output = BytesIO()
    doc.save(output)
    result = export_supervision_template(output.getvalue(), [], '2026-09-19', '张三')
    header = Document(BytesIO(result)).paragraphs[1].text
    assert re.sub(r'\s+', ' ', header) == '填写人： 张三 日期： 2026年9月19日 白班'


@pytest.mark.parametrize('change', ['title', 'weather', 'row_label', 'extra_table', 'row_count'])
def test_unrecognized_template_rejected(change):
    doc = Document(BytesIO(template()))
    if change == 'title':
        doc.paragraphs[0].text = '施工日志'
    elif change == 'weather':
        doc.tables[0].cell(0, 0).text = '错误栏'
    elif change == 'row_label':
        doc.tables[0].cell(4, 0).text = '未知栏目'
    elif change == 'extra_table':
        doc.add_table(rows=1, cols=1)
    else:
        doc.tables[0]._tbl.remove(doc.tables[0].rows[1]._tr)
    output = BytesIO()
    doc.save(output)
    with pytest.raises(ValueError):
        export_supervision_template(output.getvalue(), [], '', '')


def test_invalid_docx_rejected():
    with pytest.raises(ValueError):
        export_supervision_template(b'not a zip', [], '', '')
