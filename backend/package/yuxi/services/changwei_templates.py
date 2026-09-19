"""在已识别的监理日志原版表格中填入本次成果。"""

from copy import deepcopy
from io import BytesIO
from zipfile import BadZipFile

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE
from docx.oxml.ns import qn
from lxml.etree import XMLSyntaxError


def export_supervision_template(template_bytes, modules, period, recorder, weather=None) -> bytes:
    """填充原版监理日志，保留版式；不匹配的模板明确拒绝。"""
    try:
        document = Document(BytesIO(template_bytes))
    except (BadZipFile, KeyError, ValueError, XMLSyntaxError) as exc:
        raise ValueError('模板不是有效的 Word 文档') from exc
    def compact(text):
        return ''.join(text.split())

    if (len(document.tables) != 1 or len(document.paragraphs) < 2
            or compact(document.paragraphs[0].text) != '监理日志'
            or '填写人' not in document.paragraphs[1].text
            or '日期' not in document.paragraphs[1].text):
        raise ValueError('模板不是支持的原版监理日志')
    table = document.tables[0]
    if len(table.rows) != 7 or any(len(row.cells) != 8 for row in table.rows):
        raise ValueError('监理日志模板表格结构不匹配')
    if [compact(table.cell(0, i).text) for i in (0, 2, 4, 6)] != ['天气', '气温', '风力', '风向']:
        raise ValueError('监理日志模板天气栏不匹配')
    markers = [('施工部位', '资源投入'), ('质量', '安全'), ('巡视', '检验'), ('问题', '落实'), ('签发',), ('其他',)]
    for row, required in zip(table.rows[1:], markers):
        if (not all(word in compact(row.cells[0].text) for word in required)
                or len({cell._tc for cell in row.cells}) != 2
                or row.cells[0]._tc is row.cells[1]._tc):
            raise ValueError('监理日志模板正文栏不匹配')

    values = {item['name']: str(item.get('text') or '') for item in modules}
    _replace_paragraph(document.paragraphs[1], f'填写人：{recorder or ""}                       日期：{period or ""}')
    weather = weather or {}
    for column, chinese, english in [(1, '天气', 'weather'), (3, '气温', 'temperature'), (5, '风力', 'wind_force'), (7, '风向', 'wind_direction')]:
        value = weather.get(chinese, weather.get(english, ''))
        if column == 1 and not weather:
            value = values.get('天气信息', '')
        _replace_cell(table.cell(0, column), str(value or ''))
    groups = [
        ('施工部位及施工内容', '施工形象及资源投入'),
        ('承包人质量检验和安全作业',), ('监理检查巡视检验',),
        ('问题及处理落实',), ('监理签发意见',), ('其他事项',),
    ]
    for row, names in zip(table.rows[1:], groups):
        _replace_cell(row.cells[1], '\n'.join(values.get(name, '') for name in names if values.get(name, '')))
    for row in table.rows:
        if row.height_rule == WD_ROW_HEIGHT_RULE.EXACTLY:
            row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for element in list(row._tr.xpath('./w:trPr/w:cantSplit')):
            element.getparent().remove(element)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _replace_paragraph(paragraph, text):
    """替换全部旧内容，保留段落属性和首个文字运行的字体属性。"""
    style = next((deepcopy(run._r.rPr) for run in paragraph.runs if run._r.rPr is not None), None)
    for child in list(paragraph._p):
        if child.tag != qn('w:pPr'):
            paragraph._p.remove(child)
    run = paragraph.add_run(text)
    if style is not None:
        run._r.insert(0, style)


def _replace_cell(cell, text):
    """保留单元格属性，清除包括旧嵌套内容在内的全部填写值。"""
    paragraph = cell.paragraphs[0]
    _replace_paragraph(paragraph, text)
    for child in list(cell._tc):
        if child.tag != qn('w:tcPr') and child is not paragraph._p:
            cell._tc.remove(child)
