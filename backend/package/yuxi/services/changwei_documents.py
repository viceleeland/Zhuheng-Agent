"""长委工程资料解析和成果生成，无网络或隐式模型调用。"""

import io
import re
from pathlib import PurePosixPath
from threading import RLock
from zipfile import ZipFile

# PDFium is process-global and not thread-safe, even for separate documents.
_PDFIUM_LOCK = RLock()

MAX_FILE_BYTES = 30 * 1024 * 1024
MAX_EXPANDED_BYTES = 100 * 1024 * 1024
CATEGORIES = [
    "施工方案",
    "施工月报",
    "监理月报",
    "资金台账",
    "合同台账",
    "收发文台账",
    "检测资料",
    "规范依据",
    "其他资料",
]


def classify(filename: str) -> str:
    """给出可由用户修正的文件名分类建议。"""
    for words, category in [
        ("规范|标准|技术要求|审核依据", "规范依据"),
        ("资金|支付|工资", "资金台账"),
        ("合同", "合同台账"),
        ("来文|收文|发文", "收发文台账"),
        ("检测|检验", "检测资料"),
        ("方案", "施工方案"),
        ("监理.*月报", "监理月报"),
        ("月报", "施工月报"),
    ]:
        if re.search(words, filename):
            return category
    return "其他资料"


def unpack(filename: str, data: bytes) -> list[tuple[str, bytes]]:
    """限制 ZIP 文件数和膨胀体积，不向磁盘解压不可信路径。"""
    if len(data) > MAX_FILE_BYTES:
        raise ValueError("单文件最大 30 MB")
    if not filename.lower().endswith(".zip"):
        return [(PurePosixPath(filename.replace("\\", "/")).name, data)]
    with ZipFile(io.BytesIO(data)) as archive:
        items = [x for x in archive.infolist() if not x.is_dir()]
        if len(items) > 100 or sum(x.file_size for x in items) > MAX_EXPANDED_BYTES:
            raise ValueError("资料包最多 100 个文件、解压后最多 100 MB")
        result = []
        for item in items:
            path = PurePosixPath(item.filename.replace("\\", "/"))
            if path.is_absolute() or ".." in path.parts or item.file_size > MAX_FILE_BYTES:
                raise ValueError("资料包包含不安全路径或过大文件")
            if path.suffix.lower() == ".zip":
                raise ValueError("请先展开嵌套 ZIP")
            result.append((path.name, archive.read(item)))
        return result


def parse_document(filename: str, data: bytes) -> dict:
    """提取可核对的位置；扫描 PDF 单独标记，绝不填造识别结果。"""
    suffix = PurePosixPath(filename).suffix.lower()
    segments, warnings = [], []
    if suffix in {".docx", ".xlsx"}:
        with ZipFile(io.BytesIO(data)) as archive:
            if sum(x.file_size for x in archive.infolist()) > MAX_EXPANDED_BYTES:
                raise ValueError("Office 文件展开体积过大")
    if suffix == ".docx":
        from docx import Document

        doc = Document(io.BytesIO(data))
        for i, p in enumerate(doc.paragraphs):
            if p.text.strip():
                segments.append({"location": f"段落 {i + 1}", "text": p.text})
        for i, table in enumerate(doc.tables):
            for j, row in enumerate(table.rows):
                segments.append({"location": f"表 {i + 1} 行 {j + 1}", "text": " | ".join(c.text for c in row.cells)})
    elif suffix == ".xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
        cached = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        for sheet in wb:
            if sheet.max_row > 20000 or sheet.max_column > 200:
                warnings.append(f"{sheet.title} 声明范围较大，仅读取前 20000 行、200 列，请核对是否有遗漏")
            for row_number, (row, values) in enumerate(
                zip(
                    sheet.iter_rows(max_row=min(sheet.max_row, 20000), max_col=min(sheet.max_column, 200)),
                    cached[sheet.title].iter_rows(
                        max_row=min(sheet.max_row, 20000), max_col=min(sheet.max_column, 200)
                    ),
                ),
                1,
            ):
                cells = []
                for cell, value in zip(row, values):
                    if cell.value is None:
                        continue
                    v = value.value if cell.data_type == "f" else cell.value
                    if cell.data_type == "f" and v is None:
                        warnings.append(f"{sheet.title}!{cell.coordinate} 公式没有缓存值，需在 Excel 重算")
                        v = "[公式待重算]"
                    cells.append(f"{cell.coordinate}={v}")
                if cells:
                    segments.append({"location": f"{sheet.title}!{row_number}", "text": " | ".join(cells)})
        wb.close()
        cached.close()
    elif suffix == ".pdf":
        import pypdfium2 as pdfium

        with _PDFIUM_LOCK, pdfium.PdfDocument(data) as doc:
            for i in range(len(doc)):
                page = doc[i]
                try:
                    textpage = page.get_textpage()
                    try:
                        text = textpage.get_text_bounded().strip()
                    finally:
                        textpage.close()
                finally:
                    page.close()
                if len(text) < 20:
                    warnings.append(f"第 {i + 1} 页需要 OCR 或人工补录")
                if text:
                    segments.append({"location": f"第 {i + 1} 页", "text": text})
    elif suffix in {".txt", ".md", ".csv"}:
        segments.append({"location": "正文", "text": data.decode("utf-8-sig")})
    else:
        return {
            "status": "unsupported",
            "segments": [],
            "warnings": ["暂不支持此格式，请转换为 DOCX、XLSX、PDF 或 UTF-8 文本"],
        }
    if sum(len(s["text"]) for s in segments) > 2_000_000:
        raise ValueError("提取文本超过 200 万字，请拆分")
    bounded = []
    for segment in segments:
        for offset in range(0, len(segment["text"]), 4000):
            location = segment["location"]
            if len(segment["text"]) > 4000:
                location += f" / 字符 {offset + 1}-{min(offset + 4000, len(segment['text']))}"
            bounded.append({"location": location, "text": segment["text"][offset : offset + 4000]})
    return {
        "status": "needs_attention" if warnings or not segments else "ready",
        "segments": bounded,
        "warnings": warnings,
    }


def normalize_station(text: str) -> str:
    """统一常见口述部位和数字桩号，保留原始事实。"""
    for source, target in [
        ("左分洪洞", "FHDL"),
        ("右分洪洞", "FHDR"),
        ("左洞", "FHDL"),
        ("右洞", "FHDR"),
        ("首部明渠", "SBMQ"),
        ("尾部明渠", "WBMQ"),
    ]:
        text = text.replace(source, target)
    return re.sub(r"(\d+)\s*加\s*(\d+)", lambda m: f"{m[1]}+{m[2].zfill(3)}", text)


def export_docx(title: str, project: str, period: str, modules: list, revision: int) -> bytes:
    """生成可编辑基础版成果；客户原版模板另行映射验证。"""
    from docx import Document
    from docx.shared import Pt

    doc = Document()
    doc.styles["Normal"].font.name = "宋体"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, 0)
    doc.add_paragraph(f"工程：{project}\n报告日期 / 期间：{period}\n任务版本：{revision}")
    for module in modules:
        doc.add_heading(module["name"], 1)
        doc.add_paragraph(module["text"])
        for source in module.get("sources", []):
            doc.add_paragraph(f"资料依据：{source}")
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


def ocr_pdf_pages(data: bytes, pages: list[int]) -> list[dict]:
    """复用 Yuxi RapidOCR，在内存中按页识别，保留页码。"""
    import pypdfium2 as pdfium
    from yuxi.knowledge.parser.factory import DocumentProcessorFactory

    parser = DocumentProcessorFactory.get_processor("rapid_ocr")
    result = []
    with _PDFIUM_LOCK, pdfium.PdfDocument(data) as document:
        for number in pages:
            page = document[number - 1]
            try:
                width, height = page.get_size()
                if width * height * 2.25 > 25_000_000:
                    raise ValueError("页面尺寸过大，请先缩小 PDF 页面")
                bitmap = page.render(scale=1.5)
                try:
                    text = parser.process_image(bitmap.to_pil())
                finally:
                    bitmap.close()
            finally:
                page.close()
            result.append({"location": f"第 {number} 页", "text": text or "", "ocr": True})
    return result
